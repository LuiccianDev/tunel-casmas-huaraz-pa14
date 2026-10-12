"""Escritura de libros Excel con formato uniforme y simple."""

from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .modelo import Gantt, Libro, Tabla

_COLOR_VERDE_PRIMARIO = "2E7D32"  # Verde técnico sobrio institucional
_COLOR_VERDE_GANTT = "4CAF50"  # Verde para barras gantt
_COLOR_VERDE_TOTAL = "E8F5E9"  # Fondo verde muy suave para filas de totales
_TAMANO_FUENTE = 11  # Tamaño estricto e invariable para todo el documento

_RELLENO_ENCABEZADO = PatternFill("solid", fgColor=_COLOR_VERDE_PRIMARIO)
_RELLENO_GANTT = PatternFill("solid", fgColor=_COLOR_VERDE_GANTT)
_RELLENO_TOTAL = PatternFill("solid", fgColor=_COLOR_VERDE_TOTAL)
_FUENTE_ENCABEZADO = Font(
    name="Calibri", size=_TAMANO_FUENTE, bold=True, color="FFFFFF"
)
_FUENTE_NORMAL = Font(name="Calibri", size=_TAMANO_FUENTE)
_FUENTE_TOTAL = Font(name="Calibri", size=_TAMANO_FUENTE, bold=True, color="1B5E20")
_LINEA = Side(style="thin", color="A6A6A6")
_BORDE = Border(left=_LINEA, right=_LINEA, top=_LINEA, bottom=_LINEA)
_FILA_INICIAL = 4


def _resolver_formulas(
    celdas_formulas: list[tuple[Any, str]],
    mapa_rangos: dict[str, str],
    mapa_simbolos: dict[str, str],
) -> None:
    """Resuelve los marcadores {RANGO} y {Simbolo} en las fórmulas dinámicas."""
    rangos_ord = sorted(mapa_rangos.keys(), key=len, reverse=True)
    simbolos_ord = sorted(mapa_simbolos.keys(), key=len, reverse=True)
    for celda, formula_plantilla in celdas_formulas:
        formula_resuelta = formula_plantilla
        for r_id in rangos_ord:
            token = f"{{{r_id}}}"
            if token in formula_resuelta:
                formula_resuelta = formula_resuelta.replace(token, mapa_rangos[r_id])
        for sim in simbolos_ord:
            token = f"{{{sim}}}"
            if token in formula_resuelta:
                formula_resuelta = formula_resuelta.replace(token, mapa_simbolos[sim])
        celda.value = formula_resuelta
        if celda.number_format not in ("0.0%", "#,##0"):
            celda.number_format = "#,##0.00"


def _aplicar_validaciones(
    ws: Worksheet,
    celdas_validacion: list[tuple[Any, str]],
    mapa_rangos: dict[str, str],
) -> None:
    """Agrega listas desplegables (DataValidation) a las celdas configuradas."""
    if not celdas_validacion:
        return
    from openpyxl.worksheet.datavalidation import DataValidation

    validaciones: dict[str, DataValidation] = {}
    for celda, id_rango in celdas_validacion:
        if id_rango not in mapa_rangos:
            continue
        partes = mapa_rangos[id_rango].split(":")
        col_ini, fila_ini, fila_fin = partes[0][0], partes[0][1:], partes[1][1:]
        rango_lista = f"=${col_ini}${fila_ini}:${col_ini}${fila_fin}"
        if rango_lista not in validaciones:
            dv = DataValidation(type="list", formula1=rango_lista, allow_blank=True)
            ws.add_data_validation(dv)
            validaciones[rango_lista] = dv
        validaciones[rango_lista].add(celda)


def escribir_libro(libro: Libro, carpeta: Path, proyecto: str) -> Path:
    """Crea `carpeta/libro.archivo` con una pestaña por cada Hoja."""
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.remove(wb.active)
    for hoja in libro.hojas:
        ws = wb.create_sheet(hoja.nombre[:31])
        ws["A1"] = proyecto
        ws["A1"].font = Font(
            name="Calibri", bold=True, size=_TAMANO_FUENTE, color=_COLOR_VERDE_PRIMARIO
        )
        ws["A2"] = hoja.titulo
        ws["A2"].font = Font(
            name="Calibri", bold=True, size=_TAMANO_FUENTE, color="1B5E20"
        )
        fila = _FILA_INICIAL
        rangos_tablas = []
        mapa_simbolos: dict[str, str] = {}
        mapa_rangos: dict[str, str] = {}
        celdas_formulas: list[tuple[Any, str]] = []
        celdas_validacion: list[tuple[Any, str]] = []
        for tabla in hoja.tablas:
            fila, rango = _escribir_tabla(
                ws,
                tabla,
                fila,
                mapa_simbolos,
                mapa_rangos,
                celdas_formulas,
                celdas_validacion,
            )
            fila += 1
            rangos_tablas.append(rango)

        _resolver_formulas(celdas_formulas, mapa_rangos, mapa_simbolos)
        _aplicar_validaciones(ws, celdas_validacion, mapa_rangos)

        if len(hoja.tablas) == 1 and rangos_tablas:
            ws.auto_filter.ref = rangos_tablas[0]

        columnas_gantt = _escribir_gantt(ws, hoja.gantt, fila) if hoja.gantt else []
        _ajustar_anchos(ws)
        for letra in columnas_gantt:
            ws.column_dimensions[letra].width = 4
    carpeta.mkdir(parents=True, exist_ok=True)
    ruta = carpeta / libro.archivo
    wb.save(ruta)
    return ruta


def _formatos_columna(columnas: list[str]) -> dict[int, str]:
    """Detecta formatos sugeridos por nombre de columna."""
    claves_enteros = (
        "(pax)",
        "pax",
        "guardia",
        "portal",
        "trabajadores",
        "obreros",
        "staff",
    )
    formatos: dict[int, str] = {}
    for col, nombre in enumerate(columnas, start=1):
        nom_low = nombre.lower()
        if nombre.strip().startswith("%"):
            formatos[col] = "0.0%"
        elif any(k in nom_low for k in claves_enteros):
            formatos[col] = "#,##0"
    return formatos


def _ajustar_valor_param(
    valor_real: Any, unidad_param: str, formato_base: str | None
) -> tuple[Any, str | None]:
    """Ajusta el valor y formato para la columna Valor de una tabla_parametros."""
    if unidad_param in ("personas", "días", "dias", "und", "taladros"):
        return valor_real, "#,##0"
    if unidad_param == "%":
        if isinstance(valor_real, (int, float)) and not isinstance(valor_real, bool):
            valor_real /= 100
        return valor_real, "0.0%"
    return valor_real, formato_base


def _escribir_fila_tabla(
    ws: Worksheet,
    fila: int,
    datos: list[Any],
    es_tabla_param: bool,
    formatos_col: dict[int, str],
    mapa_simbolos: dict[str, str] | None,
    celdas_formulas: list[tuple[Any, str]] | None,
    celdas_validacion: list[tuple[Any, str]] | None,
) -> None:
    """Escribe una fila de datos de una Tabla y registra símbolos/fórmulas."""
    if es_tabla_param and mapa_simbolos is not None and len(datos) >= 2 and datos[1]:
        simbolo = str(datos[1]).strip()
        if simbolo:
            mapa_simbolos[simbolo] = f"C{fila}"

    es_total = bool(datos and str(datos[0]).strip().upper().startswith("TOTAL"))
    unidad = str(datos[3]).strip().lower() if es_tabla_param and len(datos) >= 4 else ""

    for col, valor in enumerate(datos, start=1):
        celda = ws.cell(fila, col)
        valor_real = valor
        if type(valor).__name__ == "OpcionLista":
            valor_real = valor.valor_inicial
            if celdas_validacion is not None:
                celdas_validacion.append((celda, valor.id_rango_ref))

        formato = formatos_col.get(col)
        if es_tabla_param and col == 3:
            valor_real, formato = _ajustar_valor_param(valor_real, unidad, formato)

        _celda(celda, valor_real, es_total=es_total, formato_sugerido=formato)
        if (
            celdas_formulas is not None
            and isinstance(valor_real, str)
            and valor_real.startswith("=")
            and "{" in valor_real
        ):
            celdas_formulas.append((celda, valor_real))


def _escribir_tabla(
    ws: Worksheet,
    tabla: Tabla,
    fila: int,
    mapa_simbolos: dict[str, str] | None = None,
    mapa_rangos: dict[str, str] | None = None,
    celdas_formulas: list[tuple[Any, str]] | None = None,
    celdas_validacion: list[tuple[Any, str]] | None = None,
) -> tuple[int, str]:
    if tabla.titulo:
        ws.cell(fila, 1, tabla.titulo).font = Font(
            name="Calibri", size=_TAMANO_FUENTE, bold=True, color=_COLOR_VERDE_PRIMARIO
        )
        fila += 1

    fila_encabezado = fila
    for col, nombre in enumerate(tabla.columnas, start=1):
        _encabezado(ws.cell(fila, col, nombre))

    es_tabla_param = (
        len(tabla.columnas) >= 3
        and tabla.columnas[1] == "Símbolo"
        and tabla.columnas[2] == "Valor"
    )
    formatos_col = _formatos_columna(tabla.columnas)
    fila_datos_inicio = fila + 1

    for datos in tabla.filas:
        fila += 1
        _escribir_fila_tabla(
            ws,
            fila,
            datos,
            es_tabla_param,
            formatos_col,
            mapa_simbolos,
            celdas_formulas,
            celdas_validacion,
        )

    letra_fin = get_column_letter(len(tabla.columnas))
    rango = f"A{fila_encabezado}:{letra_fin}{fila}"
    if tabla.id_rango and mapa_rangos is not None:
        mapa_rangos[tabla.id_rango] = f"A{fila_datos_inicio}:{letra_fin}{fila}"
    return fila + 1, rango


def _escribir_gantt(ws: Worksheet, gantt: Gantt, fila: int) -> list[str]:
    fin_total = max(fin for _, _, fin in gantt.actividades)
    semanas = (fin_total - gantt.inicio).days // 7 + 1
    fijas = ["Actividad", "Inicio", "Fin", "Días"]
    for col, nombre in enumerate(fijas + [f"S{i + 1}" for i in range(semanas)], 1):
        _encabezado(ws.cell(fila, col, nombre))
    for nombre, inicio, fin in gantt.actividades:
        fila += 1
        for col, valor in enumerate([nombre, inicio, fin, (fin - inicio).days + 1], 1):
            _celda(ws.cell(fila, col), valor)
        for i in range(semanas):
            sem_ini = gantt.inicio + timedelta(days=7 * i)
            celda = ws.cell(fila, len(fijas) + 1 + i)
            celda.border = _BORDE
            if inicio <= sem_ini + timedelta(days=6) and fin >= sem_ini:
                celda.fill = _RELLENO_GANTT
    return [get_column_letter(len(fijas) + 1 + i) for i in range(semanas)]


def _encabezado(celda) -> None:
    celda.fill = _RELLENO_ENCABEZADO
    celda.font = _FUENTE_ENCABEZADO
    celda.border = _BORDE
    celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _normalizar_porcentaje(valor: Any) -> tuple[Any, bool]:
    """Convierte cadenas como '15.0%' en float 0.15."""
    if (
        isinstance(valor, str)
        and valor.strip().endswith("%")
        and not valor.startswith("=")
    ):
        try:
            return float(valor.replace("%", "").strip()) / 100.0, True
        except ValueError:
            pass
    return valor, False


def _formato_por_tipo(valor: Any) -> str | None:
    """Devuelve el number_format de Excel según el tipo de dato."""
    if isinstance(valor, float):
        return "#,##0.00"
    if isinstance(valor, int):
        return "#,##0"
    if isinstance(valor, date):
        return "dd/mm/yyyy"
    if isinstance(valor, str) and valor.startswith("="):
        return "#,##0.00"
    return None


def _celda(
    celda,
    valor,
    es_total: bool = False,
    formato_sugerido: str | None = None,
) -> None:
    if isinstance(valor, bool):
        valor = "Sí" if valor else "No"

    valor, es_porcentaje = _normalizar_porcentaje(valor)
    celda.value = valor
    celda.border = _BORDE
    if es_total:
        celda.fill = _RELLENO_TOTAL
    celda.font = _FUENTE_TOTAL if es_total else _FUENTE_NORMAL

    formato = formato_sugerido or (
        "0.0%" if es_porcentaje else _formato_por_tipo(valor)
    )
    if formato:
        celda.number_format = formato
    elif isinstance(valor, str) and valor.startswith(("http://", "https://")):
        celda.hyperlink = valor
        celda.font = Font(name="Calibri", size=11, color="0563C1", underline="single")


def _ajustar_anchos(ws: Worksheet) -> None:
    anchos: dict[int, int] = {}
    for fila in ws.iter_rows(min_row=_FILA_INICIAL):
        for celda in fila:
            if celda.value is None:
                continue
            largo = (
                11 if isinstance(celda.value, date | float) else len(str(celda.value))
            )
            anchos[celda.column] = max(anchos.get(celda.column, 0), largo)
    for col, largo in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(largo + 2, 8), 60)
