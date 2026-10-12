"""Pruebas básicas: fórmulas clave y ejecución completa del flujo."""

import math
from datetime import date

import pytest
from openpyxl import load_workbook

from casmas_pa14.config import PARAMETROS, cargar_parametros
from casmas_pa14.main import ejecutar
from casmas_pa14.modulos.m01_ubicacion import distancia
from casmas_pa14.modulos.m03_seccion import area_herradura
from casmas_pa14.modulos.m03_seccion import calcular as calcular_seccion
from casmas_pa14.modulos.m09_electricidad import caudal_por_persona, grupo_comercial


def test_area_herradura():
    # B = 2, H = 2 → hw = 1 → A = 2·1 + π·4/8
    assert area_herradura(2.0, 2.0) == pytest.approx(2 + math.pi / 2)


def test_area_herradura_invalida():
    with pytest.raises(ValueError):
        area_herradura(6.0, 2.0)


def test_seccion_proyecto():
    # Plano CAD: 13.30 × 9.58 m, gálibo 11.80 × 5.50 inscrito en la bóveda
    p = cargar_parametros(PARAMETROS)
    r = calcular_seccion(p, {"ubicacion": {"longitud_m": 1000.0}})
    assert r["hastial_m"] == pytest.approx(2.93)
    assert r["area_m2"] == pytest.approx(13.30 * 2.93 + math.pi * 13.30**2 / 8)
    assert r["ancho_galibo_m"] == pytest.approx(11.80)
    assert r["holgura_galibo_m"] == pytest.approx(0.0, abs=0.01)
    assert all(f[4] == "Sí" for f in r["verificacion"])


def test_distancia():
    assert distancia({"este": 0, "norte": 0}, {"este": 3, "norte": 4}) == 5


def test_caudal_por_altitud():
    assert caudal_por_persona(1000) == 3
    assert caudal_por_persona(3500) == 5
    assert caudal_por_persona(4100) == 6


def test_grupo_comercial():
    assert grupo_comercial(480) == 500
    assert grupo_comercial(500) == 500


def test_parametros_cargan():
    p = cargar_parametros(PARAMETROS)
    assert isinstance(p["proyecto"]["fecha_inicio"], date)


def test_flujo_completo(tmp_path):
    ctx, archivos = ejecutar(PARAMETROS, tmp_path, tmp_path / "memoria")
    excels = sorted(a.name for a in archivos if a.suffix == ".xlsx")
    assert excels == [
        "01_Ubicacion.xlsx",
        "02_Geologia.xlsx",
        "03_Seccion.xlsx",
        "04_Personal.xlsx",
        "05_Campamento.xlsx",
        "06_Maquinaria.xlsx",
        "07_Polvorin.xlsx",
        "08_Agua.xlsx",
        "09_Electricidad.xlsx",
        "10_Cronograma.xlsx",
        "11_Portales.xlsx",
        "12_Voladura.xlsx",
        "13_Sostenimiento.xlsx",
        "14_Afrontonamiento.xlsx",
        "15_Voladura_Banquetas.xlsx",
    ]
    for a in archivos:
        assert a.exists() and a.stat().st_size > 0
    assert load_workbook(tmp_path / "excel" / "08_Agua.xlsx").sheetnames == [
        "Consumo",
        "PETAP",
        "PETAR",
        "Pozo",
    ]
    assert load_workbook(tmp_path / "excel" / "03_Seccion.xlsx").sheetnames == [
        "Geometria",
        "Galibo",
        "Justificacion MTC",
    ]
    hojas_maq = load_workbook(tmp_path / "excel" / "06_Maquinaria.xlsx").sheetnames
    assert hojas_maq[0:5] == [
        "Frente",
        "Accesos",
        "Banquetas",
        "Botadero",
        "Campamento",
    ]
    assert "Resumen_Flota" in hojas_maq
    assert "Sandvik DT820" in hojas_maq
    assert "Volvo FMX 6x4" in hojas_maq
    assert "Caterpillar D8" in hojas_maq
    memoria = (tmp_path / "memoria" / "03_Seccion.md").read_text(encoding="utf-8")
    assert "Verificación normativa" in memoria
    lsp = (tmp_path / "lisp" / "PA14_dibujos.lsp").read_text(encoding="ascii")
    assert lsp.count("(") == lsp.count(")"), "paréntesis LISP desbalanceados"
    assert "defun c:PA14_TODO" in lsp
    assert "defun c:PA14_GEOLOGIA" in lsp
    assert "defun c:PA14_MALLA" in lsp
    assert ctx["cronograma"]["plazo_dias"] > ctx["seccion"]["dias_excavacion"]


def test_personal_formulas_dinamicas(tmp_path):
    ejecutar(PARAMETROS, tmp_path, tmp_path / "memoria")
    wb = load_workbook(tmp_path / "excel" / "04_Personal.xlsx", data_only=False)
    ws = wb["Resumen"]

    # Fila 6: Personal de Campamento (B6 CAMP1, C6 CAMP2, D6 Dia, E6 Noc, F6 total, G6 pct)
    assert str(ws["B6"].value).startswith("=SUMIFS"), (
        "B6 desglose CAMP1 debe ser SUMIFS"
    )
    assert str(ws["C6"].value).startswith("=SUMIFS"), (
        "C6 desglose CAMP2 debe ser SUMIFS"
    )
    assert str(ws["D6"].value).startswith("=SUMIFS"), "D6 desglose Día debe ser SUMIFS"
    assert str(ws["E6"].value).startswith("=SUMIFS"), (
        "E6 desglose Noche debe ser SUMIFS"
    )
    assert str(ws["F6"].value).startswith("="), "F6 debe ser fórmula"
    assert str(ws["G6"].value).startswith("="), "G6 porcentaje debe ser fórmula"
    assert ws["G6"].number_format == "0.0%", "G6 debe tener formato de porcentaje"

    # Fila 9: Total Proyecto
    assert str(ws["B9"].value).startswith("="), "B9 debe ser fórmula"
    assert str(ws["F9"].value).startswith("="), "F9 debe ser fórmula"
    assert str(ws["G9"].value).startswith("="), "G9 debe ser fórmula"

    # Fila 13: Mano de Obra Directa (B13 Staff, C13 Obreros, D13 total, E13 pct)
    assert str(ws["B13"].value).startswith("=SUMIFS"), "B13 Staff debe ser SUMIFS"
    assert str(ws["C13"].value).startswith("=SUMIFS"), "C13 Obreros debe ser SUMIFS"
    assert str(ws["D13"].value).startswith("="), "D13 debe ser fórmula"
    assert str(ws["E13"].value).startswith("="), "E13 porcentaje debe ser fórmula"
    assert ws["E13"].number_format == "0.0%", "E13 debe tener formato de porcentaje"

    # Fila 19: Matriz Campamento x Turno (B19 C1 Día, C19 C1 Noche)
    assert str(ws["B19"].value).startswith("=SUMIFS"), "B19 C1 Día debe ser SUMIFS"
    assert str(ws["C19"].value).startswith("=SUMIFS"), "C19 C1 Noche debe ser SUMIFS"

    # Parámetros consolidados: Población (C25), Campamento (C26), Obreros (C36), Planilla (C40)
    assert str(ws["C25"].value).startswith("="), (
        "C25 Pob debe ser fórmula de suma de macro-áreas"
    )
    assert str(ws["C26"].value).startswith("=SUMIFS"), "C26 P_Camp debe ser SUMIFS"
    assert str(ws["C36"].value).startswith("="), "C36 Obreros debe ser fórmula"
    assert str(ws["C40"].value).startswith("="), "C40 Planilla debe ser fórmula"


def test_ubicacion_pendiente_se_exporta_como_porcentaje_excel(tmp_path):
    ejecutar(PARAMETROS, tmp_path, tmp_path / "memoria")
    wb = load_workbook(tmp_path / "excel" / "01_Ubicacion.xlsx", data_only=False)
    ws = wb["Eje del tunel"]

    assert ws["C5"].value.startswith("=SQRT((Coordenadas!")
    assert ws["C6"].value == "=Coordenadas!E12-Coordenadas!E11"
    assert ws["C7"].value == "=SQRT(C5^2+C6^2)"
    assert ws["C8"].value == "=C6/C5"
    assert ws["C8"].number_format == "0.0%"
    assert ws["C9"].value.startswith("=MOD(DEGREES(ATAN2(Coordenadas!")

    distancias = wb["Distancias"]
    assert distancias["C5"].value.startswith("=SQRT((Coordenadas!")
    assert distancias["E5"].value == "=C5>=D5"


def test_seccion_separa_perforacion_y_volumenes(tmp_path):
    ejecutar(PARAMETROS, tmp_path, tmp_path / "memoria")
    wb = load_workbook(tmp_path / "excel" / "03_Seccion.xlsx", data_only=False)

    geo = wb["Geometria"]
    assert geo["C9"].value == "=C7-C8"
    assert geo["C13"].value == "=C6*C11+PI()*C6^2/8"
    assert geo["C14"].value == "=C6+2*C11+PI()*C6/2"

    galibo = wb["Galibo"]
    assert galibo["C9"].value == "=C5*C6+2*C7+C8"
    assert galibo["C14"].value.startswith("=Geometria!C12+SQRT(")

    norma = wb["Justificacion MTC"]
    assert norma["B5"].value == "=Galibo!C6"
    assert norma["E5"].value == "=B5>=C5"
    assert norma["E11"].value == "=ABS(B11-C11)<=0.01"

    assert "Perforacion y avance" not in wb.sheetnames
    assert "Volumenes" not in wb.sheetnames


def test_voladura_banquetas_sostenimiento_y_afrontonamiento_separados(tmp_path):
    ctx, _ = ejecutar(PARAMETROS, tmp_path, tmp_path / "memoria")

    # 1. Verificar 12_Voladura.xlsx (exclusivo túnel subterráneo con Jumbo 3 brazos Sandvik DT1131)
    wb_vol = load_workbook(tmp_path / "excel" / "12_Voladura.xlsx", data_only=False)
    assert wb_vol.sheetnames == [
        "Malla_y_Avance_Tunel",
        "Voladura_Konya_Tunel",
        "Ficha_Jumbo_DT1131",
        "BD_Explosivos",
        "Catalogo_Detonacion",
        "BD_Rocas",
        "BD_Brocas_Equipos",
    ]
    ws_malla = wb_vol["Malla_y_Avance_Tunel"]
    assert ws_malla["C7"].value == 3
    assert str(ws_malla["C17"].value).startswith("=")
    assert (tmp_path / "memoria" / "12_Voladura.md").exists()
    assert ctx["voladura"]["frente_tunel"]["brazos_jumbo"] == 3
    assert ctx["voladura"]["frente_tunel"]["taladros_cargados"] == 231

    # 2. Verificar 15_Voladura_Banquetas.xlsx (exclusivo banquetas de portal con Epiroc PowerROC T25)
    wb_banq = load_workbook(
        tmp_path / "excel" / "15_Voladura_Banquetas.xlsx", data_only=False
    )
    assert wb_banq.sheetnames == [
        "Calculadora_Konya_T25",
        "Voladura_Por_Banqueta",
        "Ficha_PowerROC_T25",
        "BD_Explosivos",
        "BD_Rocas",
        "BD_Brocas_Equipos",
        "Criterios_y_Fuentes",
    ]
    ws_t25 = wb_banq["Calculadora_Konya_T25"]
    assert str(ws_t25["C9"].value).startswith("=IFERROR(VLOOKUP(")
    assert str(ws_t25["C13"].value).startswith("=IFERROR(VLOOKUP(")
    ws_vb = wb_banq["Voladura_Por_Banqueta"]
    assert str(ws_vb["I6"].value).startswith("=ROUNDUP(")
    assert (tmp_path / "memoria" / "15_Voladura_Banquetas.md").exists()
    assert ctx["voladura_banquetas"]["b_prac_m"] == pytest.approx(1.577, abs=0.02)
    assert len(ctx["voladura_banquetas"]["banquetas_desglose"]) == 5

    # 3. Verificar 13_Sostenimiento.xlsx (exclusivo interior del túnel) y su memoria
    wb_sost = load_workbook(
        tmp_path / "excel" / "13_Sostenimiento.xlsx", data_only=False
    )
    assert wb_sost.sheetnames == ["Sostenimiento"]
    assert (tmp_path / "memoria" / "13_Sostenimiento.md").exists()

    # 4. Verificar 14_Afrontonamiento.xlsx (afrontonamiento y sellado del portal) y su memoria
    wb_afr = load_workbook(
        tmp_path / "excel" / "14_Afrontonamiento.xlsx", data_only=False
    )
    assert wb_afr.sheetnames == [
        "RMR_Geomecanica",
        "Teoria_Cargas_Terzaghi",
        "Afrontonamiento_Portal",
        "Malla_Electrosoldada",
        "Pernos_RockBolt",
        "Shotcrete_Robotizado",
    ]
    ws_af = wb_afr["Afrontonamiento_Portal"]
    assert ws_af["C12"].value == "=Teoria_Cargas_Terzaghi!C6"
    assert ws_af["C13"].value == "=Teoria_Cargas_Terzaghi!C7"
    assert ws_af["C22"].value == "=C21-C17"
    ws_malla_af = wb_afr["Malla_Electrosoldada"]
    assert ws_malla_af["C6"].value == "=Afrontonamiento_Portal!C22"
    assert (tmp_path / "memoria" / "14_Afrontonamiento.md").exists()
    assert ctx["afrontonamiento"]["area_neta_fronton_m2"] == pytest.approx(
        493.08, abs=0.1
    )


def test_todos_los_libros_usan_formulas_dinamicas(tmp_path):
    """Verifica que los 15 libros Excel contengan fórmulas nativas y ningún {Simbolo} sin resolver."""
    ejecutar(PARAMETROS, tmp_path, tmp_path / "memoria")
    carpeta_excel = tmp_path / "excel"
    libros = sorted(carpeta_excel.glob("*.xlsx"))
    assert len(libros) == 15

    for ruta_libro in libros:
        wb = load_workbook(ruta_libro, data_only=False)
        tiene_formula = False
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for val in row:
                    if isinstance(val, str) and val.startswith("="):
                        tiene_formula = True
                        assert "{" not in val and "}" not in val, (
                            f"Marcador sin resolver en {ruta_libro.name} [{ws.title}]: {val}"
                        )
        assert tiene_formula, (
            f"El libro {ruta_libro.name} no contiene ninguna fórmula dinámica (=)"
        )
