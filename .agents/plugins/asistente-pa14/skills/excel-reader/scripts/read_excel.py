#!/usr/bin/env python3
"""Read Excel files and extract data using Python with uv."""

import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None


def read_excel_file(filepath: str, sheet_name: str = None, header_row: int = 0) -> dict:
    """Read an Excel file and extract data.

    Args:
        filepath: Path to the Excel file
        sheet_name: Specific sheet to read (None for all sheets)
        header_row: Row index to use as headers (0-indexed)

    Returns:
        Dictionary with sheet names as keys and data as values
    """
    path = Path(filepath)

    if not path.exists():
        return {"error": f"File not found: {filepath}"}

    if openpyxl is None:
        return {
            "error": "openpyxl not installed. Install with: uv pip install openpyxl"
        }

    try:
        workbook = openpyxl.load_workbook(str(path), data_only=True)

        result = {}

        sheets_to_read = []
        if sheet_name:
            if sheet_name in workbook.sheetnames:
                sheets_to_read = [workbook[sheet_name]]
        else:
            sheets_to_read = [workbook[ws] for ws in workbook.sheetnames]

        for worksheet in sheets_to_read:
            # Get column headers from specified row (1-indexed in openpyxl)
            # header_row is 0-indexed: 0 = first row, 1 = second row, etc.
            # openpyxl rows are 1-indexed, so header_row=0 -> openpyxl row 1
            header_openpyxl = header_row + 1
            headers = []
            for col in range(1, worksheet.max_column + 1):
                cell = worksheet.cell(row=header_openpyxl, column=col)
                headers.append(cell.value)

            # Get data rows (start after the header row)
            rows = []
            for row_idx in range(header_openpyxl + 1, worksheet.max_row + 1):
                row_data = {}
                for col_idx in range(1, worksheet.max_column + 1):
                    cell = worksheet.cell(row=row_idx, column=col_idx)
                    col_header = (
                        headers[col_idx - 1]
                        if col_idx - 1 < len(headers)
                        else f"col_{col_idx}"
                    )
                    if cell.value is not None:
                        row_data[col_header] = cell.value

                if row_data:
                    rows.append(row_data)

            result[worksheet.title] = {
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "rows": rows,
            }

        return result

    except Exception as e:
        return {"error": str(e)}


def main():
    """Main entry point for the script."""
    if len(sys.argv) < 2:
        print("Usage: python read_excel.py <filepath> [sheet_name] [header_row]")
        sys.exit(1)

    filepath = sys.argv[1]
    sheet_name = sys.argv[2] if len(sys.argv) > 2 else None
    header_row = int(sys.argv[3]) if len(sys.argv) > 3 else 0

    result = read_excel_file(filepath, sheet_name, header_row)
    print(json.dumps(result, indent=2, default=str))


if __name__ == "__main__":
    main()
